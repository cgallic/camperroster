"""Import original Camp Hope workbooks into a private historical roster.

Dry-run by default. --apply uses DATABASE_URL from --env-file. Does not send
messages, register current campers, infer attendance or approve health records.
Receipts contain counts and hashes only. PII never appears in command output.
"""
import argparse
import collections
import datetime
import hashlib
import io
import json
from pathlib import Path
import uuid
import warnings
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

SHEETS = {
    2022: [('Returning Campers','camper'),('New Campers','camper'),('Adult Volunteers','adult_volunteer'),('Teen Volunteers','teen_volunteer')],
    2023: [('Returning Families','camper'),('New Families','camper'),('New Adult Volunteers','adult_volunteer'),('Returning Volunteers','adult_volunteer'),('Teen Volunteers','teen_volunteer')],
    2024: [('Returning Camper Families','camper'),('New Camper Families','camper'),('New Adult Vol','adult_volunteer'),('Returning Vol','adult_volunteer'),('Teen Vol','teen_volunteer')],
    2025: [('Returning Camper Families','camper'),('New Camper Families','camper'),('New Adult Vol','adult_volunteer'),('Returning Vol','adult_volunteer'),('Teen Vol','teen_volunteer')],
    2026: [('2026 Returning Campers','camper'),('2026 New Families','camper'),('2026 Adult Volunteers','adult_volunteer'),('2026 New Adult Volunteers','adult_volunteer'),('2026 Teen Volunteers','teen_volunteer')],
}


def cell_text(value):
    if value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def extract(directory, camp_id):
    records = {}
    receipt = {'files': [], 'duplicates': 0, 'skipped_non_registration_rows': 0,
               'source_error_cells': 0, 'recovered_misformatted_cells': 0,
               'skipped_rows': [], 'rows_without_registration_id': 0, 'identity_conflicts': 0}
    for year, sheets in SHEETS.items():
        paths = list(directory.glob(f'{year}*.xlsx'))
        if len(paths) != 1:
            raise ValueError(f'Expected one workbook for {year}, found {len(paths)}')
        path = paths[0]
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        with warnings.catch_warnings(record=True) as source_warnings:
            warnings.simplefilter('always')
            book = openpyxl.load_workbook(io.BytesIO(path.read_bytes()), read_only=True, data_only=True)
            sheet_counts = {}
            for sheet_name, kind in sheets:
                sheet = book[sheet_name]
                # Some original phone cells are incorrectly date-formatted.
                # openpyxl replaces out-of-range dates with #VALUE!; recover the
                # original numeric cell text from XLSX rather than lose a phone.
                with zipfile.ZipFile(path) as archive:
                    xml = ET.fromstring(archive.read(sheet._worksheet_path))
                ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                raw_numbers = {c.attrib['r']:c.find('s:v',ns).text for c in xml.findall('.//s:c',ns)
                               if c.attrib.get('t','n')=='n' and c.find('s:v',ns) is not None}
                rows = sheet.iter_rows(values_only=True)
                header = None
                header_row = 0
                for row_number, row in enumerate(rows, 1):
                    if row_number > 8:
                        raise ValueError(f'Header missing in {year}/{sheet_name}')
                    if 'First Name' in row and 'Last Name' in row:
                        header, header_row = row, row_number
                        break
                if header is None:
                    raise ValueError(f'Empty required sheet: {year}/{sheet_name}')
                # Repeated headers such as Paid are preserved with column suffixes.
                frequencies = collections.Counter(cell_text(v) for v in header if v is not None)
                keys = [f'{cell_text(v)} [column {i+1}]' if v is not None and frequencies[cell_text(v)] > 1
                        else cell_text(v) or f'[column {i+1}]' for i,v in enumerate(header)]
                count = 0
                for row_number,row in enumerate(rows,header_row+1):
                    if not any(v is not None for v in row):
                        continue
                    data = {keys[i] if i < len(keys) else f'[column {i+1}]': cell_text(v)
                            for i,v in enumerate(row) if v is not None}
                    for i,v in enumerate(row):
                        if v == '#VALUE!':
                            cell = f'{openpyxl.utils.get_column_letter(i+1)}{row_number}'
                            if cell in raw_numbers:
                                field = keys[i] if i < len(keys) else f'[column {i+1}]'
                                data[field] = raw_numbers[cell]
                                data.setdefault('_recovered_original_numeric_cells',[]).append(cell)
                                receipt['recovered_misformatted_cells'] += 1
                    first, last = data.get('First Name'), data.get('Last Name')
                    if not first or not last or first == 'First Name':
                        receipt['skipped_non_registration_rows'] += 1
                        receipt['skipped_rows'].append({'year':year,'sheet':sheet_name,'row':row_number,'reason':'missing participant name or repeated header'})
                        continue
                    registration_id = data.get('Registration ID')
                    if year < 2026 and not registration_id:
                        receipt['rows_without_registration_id'] += 1
                    # Source IDs distinguish registrations; 2026 identity is per
                    # person/year/type. DOB/email only disambiguate, never inferred.
                    identity = registration_id or '|'.join([
                        first.casefold(),last.casefold(),data.get('Birth Date') or data.get('Date of Birth') or '',
                        (data.get('Parent Email') or data.get('Email') or data.get('Email Address') or '').casefold()])
                    key = str(uuid.uuid5(uuid.NAMESPACE_URL, f'camperroster:{camp_id}:{year}:{kind}:{identity}'))
                    rec = dict(id=key,camp_id=camp_id,season_year=year,participant_type=kind,
                               source_registration_id=registration_id,first_name=first,last_name=last,
                               source_status=data.get('Registration Status'),source_workbook=path.name,
                               source_sheet=sheet_name,source_row=row_number,source_sha256=digest,
                               source_data=data)
                    count += 1
                    if key in records:
                        previous = records[key]
                        if (previous['first_name'].casefold(),previous['last_name'].casefold()) != (first.casefold(),last.casefold()):
                            receipt['identity_conflicts'] += 1
                            raise ValueError(f'Conflicting source registration ID in {year}/{sheet_name} row {row_number}')
                        previous['source_data'].setdefault('_additional_source_rows',[]).append(
                            {'sheet':sheet_name,'row':row_number,'fields':data})
                        receipt['duplicates'] += 1
                    else:
                        records[key] = rec
                sheet_counts[sheet_name] = count
            book.close()
            receipt['source_error_cells'] += len(source_warnings)
        receipt['files'].append({'year':year,'file':path.name,'sha256':digest,'sheets':sheet_counts})
    receipt['counts'] = {str(year):dict(collections.Counter(r['participant_type'] for r in records.values() if r['season_year']==year)) for year in SHEETS}
    receipt['total'] = len(records)
    return list(records.values()), receipt


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-dir',type=Path,required=True)
    parser.add_argument('--camp-id',required=True)
    parser.add_argument('--receipt',type=Path,required=True)
    parser.add_argument('--env-file',type=Path)
    parser.add_argument('--apply',action='store_true')
    args=parser.parse_args()
    records,receipt=extract(args.source_dir,args.camp_id)
    if args.apply:
        import psycopg
        from psycopg.types.json import Jsonb
        env={k:v.strip().strip('"') for line in args.env_file.read_text().splitlines()
             if '=' in line and not line.startswith('#') for k,v in [line.split('=',1)]}
        with psycopg.connect(env['DATABASE_URL'],connect_timeout=20) as conn:
            camp=conn.execute('select name from public.camps where id=%s',(args.camp_id,)).fetchone()
            if not camp or camp[0] != 'Camp Hope':
                raise ValueError('Target is not the confirmed Camp Hope tenant')
            existing=dict(conn.execute('select id::text,source_sha256 from public.camp_registration_imports where camp_id=%s',(args.camp_id,)).fetchall())
            if any(r['id'] in existing and existing[r['id']] != r['source_sha256'] for r in records):
                raise ValueError('Existing import uses changed workbook; reconcile explicitly before replacing')
            columns=list(records[0])
            sql='insert into public.camp_registration_imports ('+','.join(columns)+') values ('+','.join(['%s']*len(columns))+') on conflict (id) do nothing'
            with conn.cursor() as cur:
                cur.executemany(sql,[[Jsonb(r[c]) if c=='source_data' else r[c] for c in columns] for r in records])
            found={r[0] for r in conn.execute('select id::text from public.camp_registration_imports where camp_id=%s',(args.camp_id,))}
            assert all(r['id'] in found for r in records), 'Readback missing source rows'
            receipt['inserted']=sum(r['id'] not in existing for r in records)
            receipt['already_present']=sum(r['id'] in existing for r in records)
            receipt['verified_ids']=len(records)
            receipt['state']='imported'
    else:
        receipt['state']='validated_dry_run'
    args.receipt.parent.mkdir(parents=True,exist_ok=True)
    args.receipt.write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps(receipt,indent=2))


if __name__=='__main__':
    main()
